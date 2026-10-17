"""Fonctions de base du lecteur PDF : édition, conversion Word/Excel, OCR.

Utilisable seul en ligne de commande :
    python pdftools.py word   document.pdf  [sortie.docx]
    python pdftools.py excel  document.pdf  [sortie.xlsx]
    python pdftools.py ocr    scan.pdf      [sortie.pdf]   (PDF image -> PDF avec texte)
    python pdftools.py texte  document.pdf  [sortie.txt]
    python pdftools.py remplacer document.pdf "ancien" "nouveau" [sortie.pdf]
"""
import os
import re
import shutil
import subprocess
import sys
import tempfile

import pymupdf
import pdfplumber
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------- édition

def _hex_to_rgb(color):
    color = (color or "#000000").lstrip("#")
    return tuple(int(color[i:i + 2], 16) / 255 for i in (0, 2, 4))


def _int_to_rgb(value):
    return ((value >> 16) & 255) / 255, ((value >> 8) & 255) / 255, (value & 255) / 255


FONT = pymupdf.Font("helv")  # Helvetica avec €, œ, accents...


def _write(page, point, text, size, color):
    """Écrit une ligne ; `point` est la ligne de base en coordonnées affichées."""
    pt = pymupdf.Point(point) * page.derotation_matrix
    tw = pymupdf.TextWriter(page.rect)
    tw.append(pt, text, font=FONT, fontsize=size)
    morph = (pt, pymupdf.Matrix(page.rotation)) if page.rotation else None
    tw.write_text(page, color=color, morph=morph)


def add_text(doc, page_no, x, y, text, size=12, color="#000000"):
    """Écrit `text` sur la page, (x, y) étant le coin haut-gauche tel qu'affiché."""
    page = doc[page_no]
    for i, line in enumerate(text.splitlines() or [""]):
        _write(page, (x, y + size * (0.85 + i * 1.2)), line, size, _hex_to_rgb(color))


def whiteout(doc, page_no, x0, y0, x1, y1):
    """Efface définitivement une zone (texte et images) et la remplit de blanc."""
    page = doc[page_no]
    rect = pymupdf.Rect(x0, y0, x1, y1) * page.derotation_matrix
    page.add_redact_annot(rect.normalize(), fill=(1, 1, 1))
    page.apply_redactions()


def _span_style(page, rect):
    """Taille et couleur du texte d'origine situé dans `rect`."""
    for block in page.get_text("dict", clip=rect)["blocks"]:
        for line in block.get("lines", []):
            for span in line["spans"]:
                if span["text"].strip():
                    return span["size"], _int_to_rgb(span["color"]), span["origin"]
    return rect.height * 0.8, (0, 0, 0), (rect.x0, rect.y1 - rect.height * 0.2)


def replace_text(doc, old, new, pages=None):
    """Remplace toutes les occurrences de `old` par `new`. Renvoie le nombre de remplacements."""
    count = 0
    for page in doc:
        if pages is not None and page.number not in pages:
            continue
        hits = page.search_for(old)
        if not hits:
            continue
        styles = [_span_style(page, r) for r in hits]
        for r in hits:
            page.add_redact_annot(r, fill=False)
        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE,
                              graphics=pymupdf.PDF_REDACT_LINE_ART_NONE)
        for r, (size, color, origin) in zip(hits, styles):
            tw = pymupdf.TextWriter(page.rect)
            tw.append((r.x0, origin[1]), new, font=FONT, fontsize=size)
            tw.write_text(page, color=color)
            count += 1
    return count


def rotate_page(doc, page_no, angle=90):
    page = doc[page_no]
    page.set_rotation((page.rotation + angle) % 360)


def delete_page(doc, page_no):
    if len(doc) > 1:
        doc.delete_page(page_no)


# ---------------------------------------------------------------- texte / OCR

def needs_ocr(doc):
    """Vrai si le PDF ne contient presque pas de texte (scan, photo)."""
    chars = sum(len(p.get_text().strip()) for p in doc)
    return chars < 20 * len(doc)


def find_tesseract():
    """Chemin de Tesseract (logiciel libre de reconnaissance de texte)."""
    found = shutil.which("tesseract")
    if found:
        return found
    for path in (r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                 r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
                 "/opt/homebrew/bin/tesseract", "/usr/local/bin/tesseract"):
        if os.path.exists(path):
            return path
    raise RuntimeError("Tesseract n'est pas installé")


def ocr_pdf(src_path, dst_path, language="fra+eng", dpi=300):
    """Transforme un PDF image en PDF dont le texte est sélectionnable et recherchable.

    L'image de chaque page est conservée ; Tesseract ajoute dessous une couche
    de texte invisible, alignée sur les mots.
    """
    tesseract = find_tesseract()
    src = pymupdf.open(src_path)
    out = pymupdf.open()
    with tempfile.TemporaryDirectory() as tmp:
        for page in src:
            if len(page.get_text().strip()) > 20:  # déjà du texte : on garde la page
                out.insert_pdf(src, from_page=page.number, to_page=page.number)
                continue
            png = os.path.join(tmp, f"p{page.number}.png")
            pix = page.get_pixmap(dpi=dpi)
            pix.set_dpi(dpi, dpi)
            pix.save(png)
            base = os.path.join(tmp, f"p{page.number}")
            subprocess.run([tesseract, png, base, "-l", language, "pdf"], check=True,
                           capture_output=True)
            with pymupdf.open(base + ".pdf") as ocr_page:
                new = out.new_page(width=page.rect.width, height=page.rect.height)
                new.show_pdf_page(new.rect, ocr_page, 0)
    out.save(dst_path, garbage=3, deflate=True)
    return dst_path


def extract_text(pdf_path):
    with pymupdf.open(pdf_path) as doc:
        return "\n\n".join(f"--- Page {p.number + 1} ---\n{p.get_text()}" for p in doc)


# ---------------------------------------------------------------- conversions

def to_word(pdf_path, docx_path):
    from pdf2docx import Converter
    cv = Converter(pdf_path)
    try:
        cv.convert(docx_path)
    finally:
        cv.close()
    return docx_path


def _clean(cell):
    if cell is None:
        return None
    cell = str(cell).replace("\n", " ").strip()
    # Nombres au format français ("1 234,56" ou "12,5 %") -> vrais nombres Excel.
    num = cell.replace(" ", "").replace(" ", "").replace("€", "").replace("%", "")
    if re.fullmatch(r"-?\d+(,\d+)?", num):
        return float(num.replace(",", ".")) if "," in num else int(num)
    if re.fullmatch(r"-?\d+(\.\d+)?", num):
        return float(num) if "." in num else int(num)
    return cell


def _autosize(ws):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 2, 60)


def to_excel(pdf_path, xlsx_path):
    """Chaque tableau détecté devient une feuille. Sans tableau, le texte est mis en colonnes."""
    wb = Workbook()
    wb.remove(wb.active)
    n_tables = 0
    with pdfplumber.open(pdf_path) as pdf:
        for p_no, page in enumerate(pdf.pages, 1):
            tables = page.extract_tables()
            if not tables:
                tables = page.extract_tables({"vertical_strategy": "text",
                                              "horizontal_strategy": "text"})
                tables = [t for t in tables if t and len(t[0]) > 1 and len(t) > 1]
            for t_no, table in enumerate(tables, 1):
                n_tables += 1
                ws = wb.create_sheet(f"Page{p_no}-Tableau{t_no}"[:31])
                for row in table:
                    cells = [_clean(c) for c in row]
                    if any(c not in (None, "") for c in cells):
                        ws.append(cells)
                _autosize(ws)
        if n_tables == 0:
            ws = wb.create_sheet("Texte")
            for p_no, page in enumerate(pdf.pages, 1):
                for line in (page.extract_text() or "").splitlines():
                    ws.append([p_no] + [_clean(c) for c in re.split(r"\s{2,}", line)])
            _autosize(ws)
    wb.save(xlsx_path)
    return n_tables


# ---------------------------------------------------------------- ligne de commande

def _main(argv):
    if len(argv) < 2:
        print(__doc__)
        return 1
    cmd, src = argv[0], argv[1]
    base = os.path.splitext(src)[0]
    if cmd == "word":
        print(to_word(src, argv[2] if len(argv) > 2 else base + ".docx"))
    elif cmd == "excel":
        dst = argv[2] if len(argv) > 2 else base + ".xlsx"
        print(f"{to_excel(src, dst)} tableau(x) -> {dst}")
    elif cmd == "ocr":
        print(ocr_pdf(src, argv[2] if len(argv) > 2 else base + "-ocr.pdf"))
    elif cmd == "texte":
        dst = argv[2] if len(argv) > 2 else base + ".txt"
        with open(dst, "w", encoding="utf-8") as f:
            f.write(extract_text(src))
        print(dst)
    elif cmd == "remplacer":
        dst = argv[4] if len(argv) > 4 else base + "-modifie.pdf"
        with pymupdf.open(src) as doc:
            n = replace_text(doc, argv[2], argv[3])
            doc.save(dst, garbage=3, deflate=True)
        print(f"{n} remplacement(s) -> {dst}")
    else:
        print(__doc__)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(_main(sys.argv[1:]))
